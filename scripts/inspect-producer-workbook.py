import collections
import json
import pathlib
import re
import sys
import unicodedata

import openpyxl


def key(value):
    value = unicodedata.normalize('NFKD', str(value or ''))
    return re.sub(r'\s+', ' ', ''.join(c for c in value if not unicodedata.combining(c))).strip().upper()


def cpf_valid(value):
    digits = re.sub(r'\D', '', str(value or ''))
    if len(digits) != 11 or len(set(digits)) == 1:
        return False
    for length in (9, 10):
        check = (sum(int(digits[i]) * (length + 1 - i) for i in range(length)) * 10) % 11
        if (0 if check == 10 else check) != int(digits[length]):
            return False
    return True


workbook = openpyxl.load_workbook(sys.argv[1], read_only=True, data_only=True)
summary = {}
for sheet in workbook:
    values = list(sheet.values)
    summary[sheet.title] = {'rows': sheet.max_row, 'columns': sheet.max_column}
    if sheet.title in ('CONSOLIDADO', 'PRODUTORES PAF VNA', 'PRODUTORES INDEPENDENTES'):
        rows = [row for row in values[1:] if row[1] and isinstance(row[1], str)]
        name_index, area_index, tech_index = (1, 6, 7)
        summary[sheet.title].update({
            'records': len(rows),
            'uniqueNames': len(set(key(row[name_index]) for row in rows)),
            'totalArea': round(sum(float(row[area_index] or 0) for row in rows if isinstance(row[area_index], (int, float))), 3),
            'technicians': dict(collections.Counter(key(row[tech_index]) for row in rows)),
            'repeatedNames': dict((name, count) for name, count in collections.Counter(key(row[name_index]) for row in rows).items() if count > 1),
        })
    if sheet.title in ('COOPAFAMITA', 'ANTONIO ROCHA'):
        cpf_index = 2 if sheet.title == 'COOPAFAMITA' else 5
        rows = [row for row in values[1:] if isinstance(row[0], (int, float)) and row[1]]
        documents = [re.sub(r'\D', '', str(row[cpf_index] or '')) for row in rows]
        summary[sheet.title].update({
            'records': len(rows),
            'cpfsPresent': sum(bool(value) for value in documents),
            'invalidCpfs': [row[1].strip() for row in rows if row[cpf_index] and not cpf_valid(row[cpf_index])],
            'repeatedCpfCount': sum(count > 1 for doc, count in collections.Counter(doc for doc in documents if doc).items()),
            'technicians': dict(collections.Counter(key(row[7 if sheet.title == 'COOPAFAMITA' else 14]) for row in rows)),
        })
print(json.dumps(summary, ensure_ascii=True, indent=2))
